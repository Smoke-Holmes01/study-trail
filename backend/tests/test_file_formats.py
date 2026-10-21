import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from study_trail.files import extract
from study_trail.core import Problem


def test_text_pdf_scanned_and_encrypted(tmp_path):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 700 Td (Matrix dimension lesson) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    target = tmp_path / "text.pdf"
    writer.write(target)
    text, metadata = extract(target, ".pdf")
    assert "Matrix dimension lesson" in text
    assert metadata["boundaries"][0]["page"] == 1
    scanned = PdfWriter()
    scanned.add_blank_page(width=612, height=792)
    target = tmp_path / "scan.pdf"
    scanned.write(target)
    with pytest.raises(Problem) as error:
        extract(target, ".pdf")
    assert error.value.code == "NO_READABLE_TEXT"
    writer.encrypt("test-password")
    target = tmp_path / "encrypted.pdf"
    writer.write(target)
    with pytest.raises(Problem) as error:
        extract(target, ".pdf")
    assert error.value.code == "FILE_PARSE_FAILED"


def test_docx_paragraph_and_table_text(tmp_path):
    doc = Document()
    doc.add_paragraph("知识点：矩阵维度。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "左矩阵"
    table.cell(0, 1).text = "右矩阵"
    table.cell(1, 0).text = "2×3"
    table.cell(1, 1).text = "3×4"
    target = tmp_path / "lesson.docx"
    doc.save(target)
    text, metadata = extract(target, ".docx")
    assert "知识点：矩阵维度。" in text and "2×3" in text and "3×4" in text
    assert any(b["kind"] == "table" for b in metadata["boundaries"])
